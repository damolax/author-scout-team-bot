"""Bounded, side-effect-free workbook parsing for Letterdesk."""
import csv
import io
import re
import zipfile
from itertools import islice
from openpyxl import load_workbook

ALIASES = {
 'id':['Author Scout ID','Canonical Author ID','Source Row ID'],
 'name':['Author','Author Name','Name','Author Name — Verified','Author Name — Bot'],
 'email':['Contact Email','Email','Verified Email','Public Contact Email','Public Email','Public Professional Email','Verified Public Email','Author Email','Public Professional Email — Bot'],
 'website':['Website / Profile','Website','Official Website','Verified Official Website','Author Website'],
 'email_source_url':['Email Source','Email Source URL','Exact Email Source URL','Contact Source'],
 'country':['Country','Country Verified','Country / Market'],
 'genre':['Genre','Genre / Category'],
 'subject':['Selected Subject','Subject','Subject Line','Subject Option 1'],
 'body':['First Outreach Message','Best First Message — Author Language','First Outreach Message (Icelandic)','First Message (Author Language)','Personalized Message (Author Language)','First Message (Icelandic)','Icelandic Message','Agent-Style First Message','Full First Message','Personalized First Message','First Message','Message','Body','Email Body'],
 'body_english':['Best First Message — English','English Version','First Outreach Message (English)','First Message (English)','Personalized First Message (English)','English Message'],
 'bio':['Short Bio','Research Summary','Verified Research Summary','Bio','Biography','Researched Bio'],
 'books':['Key Works','Books','Main Books','Key Work Referenced','Book/Work Referenced','Books / Works'],
 'recent_activity':['Recent Activity / Current Moment','Recent Activity','Current Moment','2026 Activity Evidence'],
 'status':['Outreach Stage','Outreach Status','Processing Status','Research Status','Status'],
 'reply_notes':['Reply / Conversation Notes','Conversation Notes','Research Notes'],
}

def norm(value):
 return re.sub(r'[^a-z0-9]+',' ',str(value or '').lower()).strip()

def parse_upload(data, filename):
 if len(data)>15*1024*1024: raise ValueError('File exceeds 15 MB.')
 tables=[]
 if filename.lower().endswith('.csv'):
  try: tables=[('CSV',list(islice(csv.reader(io.StringIO(data.decode('utf-8-sig'))),5022)))]
  except UnicodeError: raise ValueError('Save the CSV using UTF-8 encoding.')
 elif filename.lower().endswith(('.xlsx','.xlsm')):
  try:
   with zipfile.ZipFile(io.BytesIO(data)) as z:
    if sum(f.file_size for f in z.infolist())>80*1024*1024: raise ValueError('Workbook expands beyond 80 MB. Split it into smaller files.')
   book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
   try:
    for sheet in book.worksheets:
     if sheet.max_row and sheet.max_row>5020: raise ValueError('Use at most 5,000 rows per sheet.')
     if sheet.max_column and sheet.max_column>150: raise ValueError('Use at most 150 columns per sheet.')
     tables.append((sheet.title,list(islice(sheet.iter_rows(values_only=True),5022))))
   finally: book.close()
  except ValueError: raise
  except Exception: raise ValueError('Could not read this Excel file. Use XLSX, XLSM or CSV.')
 else: raise ValueError('Upload an XLSX, XLSM or CSV file. Save older XLS files as XLSX first.')
 for title,table in tables:
  if not table: continue
  columns=None; header_row=0
  for header_row,raw_header in enumerate(table[:20]):
   header=[norm(x) for x in raw_header]
   candidate={k:next((header.index(norm(a)) for a in names if norm(a) in header),None) for k,names in ALIASES.items()}
   if candidate['name'] is not None and candidate['email'] is not None and (candidate['body'] is not None or candidate['body_english'] is not None):
    columns=candidate;break
  if columns is None:continue
  if len(table)-header_row>5001: raise ValueError('Use at most 5,000 rows per upload.')
  rows=[]; issues=[]; seen=set()
  for number,raw in enumerate(table[header_row+1:],header_row+2):
   if not any(v is not None and str(v).strip() for v in raw):continue
   item={k:(str(raw[i]).strip() if i is not None and i<len(raw) and raw[i] is not None else '') for k,i in columns.items()}
   item['email']=item['email'].lower()
   item['body']=item['body'] or item['body_english']
   item['row']=number
   reason=''
   if not item['name']: reason='Missing author name'
   elif not re.fullmatch(r'[^\s@<>;,]+@[^\s@<>;,]+\.[^\s@<>;,]+',item['email']) or len(item['email'])>254: reason='Missing or invalid email'
   elif not item['subject'] or not item['body']: reason='Missing subject or first message'
   elif '\n' in item['subject'] or '\r' in item['subject']: reason='Subject must be a single line'
   elif len(item['subject'])>400 or any(len(item[k])>30000 for k in ('body','body_english','bio','books','reply_notes')):reason='Text exceeds field limit'
   elif any(norm(item['status']).startswith(x) for x in ('sent','replied','on hold','hold','do not contact','unsubscribed','deprioritized','deprioritised')): reason='Excluded by outreach status: '+item['status']
   elif item['email'] in seen: reason='Duplicate email in this workbook'
   if reason: issues.append({'row':number,'name':item['name'],'reason':reason});continue
   seen.add(item['email']);rows.append(item)
  return {'sheet':title,'rows':rows,'issues':issues,'columns':{k:str(table[header_row][i]) for k,i in columns.items() if i is not None}}
 raise ValueError('No author sheet found. Include Author, Email, Subject and First Message (or First Outreach Message (English)) columns.')
