"""Letter Desk Phase 4 file parser: six services, optional messages, source fields.

Pure parser. Does not import, send, search the web, or alter the database.
Accepts CSV/XLSX/XLSM within conservative size/decompression and row limits.
"""
from __future__ import annotations
import csv
import io
import re
import zipfile
from itertools import islice
from urllib.parse import urlparse

from openpyxl import load_workbook

MAX_BYTES = 64 * 1024 * 1024
MAX_UNCOMPRESSED = 192 * 1024 * 1024
MAX_ROWS = 5000
MAX_BULK_ROWS = 50000
MAX_COLUMNS = 150
MAX_CELL_LENGTH = 30000

COMMON = {
 "name": ["Name", "Lead Name", "Contact Name"],
 "email": ["Email", "Email Address", "Contact Email", "Public Professional Email", "Verified Email", "Verified Public Email", "Author Email", "To", "Recipient Email", "Public Email"],
 "website": ["Website", "Website URL", "Official Website", "Author Website", "Business Website", "Website / Profile"],
 "maps_url": ["Google Maps URL", "Maps URL", "Maps Link"],
 "country": ["Country", "Country Verified", "Country / Market"],
 "city": ["City", "Location"],
 "source_url": ["Source URL", "Discovery Source", "Source Link", "Email Source", "Email Source URL", "Exact Email Source URL", "Contact Source"],
 "evidence": ["Evidence Notes", "Research Notes", "Research Summary", "Verified Research Summary", "Recent Activity", "Discovery Evidence", "Notes"],
 "subject": ["Subject", "Subject Line", "Selected Subject", "Personalized Subject", "Final Subject", "Message 01 Subject"],
 "message": ["First Message", "First Outreach Message", "First Message (Author Language)", "Personalized Message (Author Language)", "Best First Message — Author Language", "Message", "Email Body", "Body", "Primary Message", "Agent-Style First Message"],
 "english": ["First Message English", "First Message (English)", "English Version", "English Review", "Best First Message — English"],
 "do_not_contact": ["Do Not Contact", "DNC", "Opt Out", "Opted Out", "Unsubscribed"],
 "status": ["Status", "Outreach Status", "Outreach Stage", "Processing Status", "Research Status", "Send Gate"],
}
SERVICES = {
 "author": {
   "name": ["Author", "Author Name", "Writer", "Author Name — Verified", "Author Name — Bot"],
   "book_titles": ["Book Titles","Books","Main Books","Key Works","Key Work Referenced","Books / Works"],
   "genre": ["Genre","Genre / Category"],"recent_publication":["Recent Publication","Publication Date"],
   "publisher":["Publisher"],"author_profile":["Author Profile","Author LinkedIn"],
   "recent_activity":["Recent Activity / Current Moment","Current Moment","Recent Activity"],
   "literary_angle":["Literary Angle"],"email_source_url":["Email Source URL"],
 },
 "automation": {"name":["Business Name","Company Name","Company"],"business_type":["Business Type","Industry"],"current_tools":["Current Tools","Tech Stack"],"workflow_signal":["Manual Workflow Signal","Workflow Problem"],"integration_opportunity":["Integration Opportunity","Automation Need"],"opportunity_evidence_url":["Opportunity Evidence URL"]},
 "website-design":{"name":["Business Name","Company"],"website_issue":["Website Issue","Web Issue"],"issue_page_url":["Page URL","Issue Page URL"],"ux_observation":["UX/Accessibility Observation","UX Observation","Accessibility Issue"],"website_technology":["Website Technology","CMS"],"opportunity_evidence_url":["Evidence URL"]},
 "digital-planner":{"name":["Creator or Shop Name","Creator Name","Shop Name"],"storefront_url":["Storefront URL","Shop URL"],"planner_niche":["Planner Niche","Niche"],"product_type":["Product Type"],"product_url":["Product URL"],"audience":["Audience"],"opportunity_evidence_url":["Opportunity Evidence URL"]},
 "shopify":{"name":["Store Name","Shop Name"],"shopify_store_url":["Shopify Store URL","Store URL"],"product_category":["Product Category","Store Niche"],"store_issue":["Store Issue","Checkout Issue"],"issue_page_url":["Issue Page URL","Product URL"],"ecommerce_opportunity":["Ecommerce Opportunity"],"opportunity_evidence_url":["Evidence URL"]},
 "google-maps":{"name":["Business Name","Place Name"],"maps_url":["Google Maps URL","Maps URL","Map Link"],"place_id":["Google Place ID","Place ID"],"business_category":["Business Category","Category"],"address":["Address"],"business_website":["Business Website"],"listing_observation":["Listing Observation","Google Business Profile Observation"],"opportunity_evidence_url":["Evidence URL"]},
}
MESSAGE_COLUMN=re.compile(r"^message\s*(\d{1,8})\s*(yes|no|question|objection|later|noreply|review|optout|bounce)?$",re.I)
BRANCH_SUFFIX={"yes":"Yes","no":"No","question":"Question","objection":"Objection","later":"Later",
              "noreply":"NoReply","review":"Review","optout":"OptOut","bounce":"Bounce"}
EMAIL_RE=re.compile(r"^[^@\s<>;,]{1,100}@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+$",re.I)
TRUE_VALUES={"1","true","yes","y","opted out","unsubscribed","unsubscribe","do not contact"}
EXCLUDED_STATUSES={"sent","replied","deleted","hold","on hold","opt out","unsubscribed","do not contact"}

def norm(value):
    return re.sub(r"[^a-z0-9]+"," ",str(value or "").lower()).strip()

def clean_cell(value):
    if value is None:return ""
    if isinstance(value,(int,float,bool)):return str(value)
    s=str(value).replace("\x00","").strip()
    return s[:MAX_CELL_LENGTH+1]

def public_link(raw):
    value=str(raw or "").strip()[:1800]
    if not value:return ""
    try:
        p=urlparse(value)
        if p.scheme not in ("http","https") or not p.hostname or p.username or p.password:return ""
        if p.hostname.lower() in ("localhost","127.0.0.1") or p.hostname.endswith((".local",".internal")):return ""
        return value
    except ValueError:return ""

def _map_header(header,service):
    idx={norm(label):i for i,label in enumerate(header) if norm(label)}
    columns={}
    # Service names take precedence, but generic Name must always work too.
    all_aliases={key:list(aliases) for key,aliases in COMMON.items()}
    for key,aliases in SERVICES[service].items():
        all_aliases[key]=list(aliases)+all_aliases.get(key,[])
    for key,aliases in all_aliases.items():
        value=next((idx[norm(alias)] for alias in aliases if norm(alias) in idx),None)
        if value is not None:columns[key]=value
    branches={}
    for label,i in ((str(x or ""),i) for i,x in enumerate(header)):
        compact=norm(label).replace(" ","")
        match=MESSAGE_COLUMN.fullmatch(compact)
        if match:
            number,suffix=match.groups()
            code="Message"+number+(BRANCH_SUFFIX.get((suffix or "").lower(),""))
            if code not in branches:branches[code]=i
    return columns,branches

def _find_sheet(rows,service):
    for i,row in enumerate(rows[:20]):
        cols,branches=_map_header(row,service)
        if "name" in cols or "email" in cols or "website" in cols or "maps_url" in cols or "shopify_store_url" in cols:
            return i,cols,branches
    return None,None,None

def _process_table(name,header_idx,header,columns,branches,iterator,service,*,first_row=None,seen=None):
    items=[];errors=[]
    if seen is None:seen=set()
    for number,raw in enumerate(iterator,start=first_row if first_row is not None else header_idx+2):
        if not any(clean_cell(x) for x in raw):continue
        if len(items)+len(errors)>=MAX_ROWS:raise ValueError("Maximum 5,000 nonempty records in a standard import; use Bulk upload for up to 50,000.")
        def read(key):
            i=columns.get(key)
            return clean_cell(raw[i]) if i is not None and i<len(raw) else ""
        record={key:read(key) for key in set(COMMON)|set(SERVICES[service])}
        person=record.get("name","").strip()
        website=record.get("website") or record.get("shopify_store_url") or record.get("storefront_url") or record.get("business_website") or ""
        maps_url=record.get("maps_url") or ""
        place_id=record.get("place_id") or ""
        email=record["email"].lower().strip()
        msg=record["message"] or ""
        branch_data={}
        for code,i in branches.items():
            if i<len(raw):
                branch=clean_cell(raw[i])
                if branch:branch_data[code]=branch
        if not msg:
            msg=branch_data.get("Message01") or ""
        reason=[]
        if not person:reason.append("Missing prospect name")
        if not (email or website or maps_url or place_id):reason.append("Provide an email, official website, Google Maps URL, or Place ID")
        if email and (len(email)>254 or not EMAIL_RE.fullmatch(email)):reason.append("Email address has invalid format")
        if website and not public_link(website):reason.append("Website URL must use public http:// or https://")
        if maps_url and not public_link(maps_url):reason.append("Maps URL must use public http:// or https://")
        if record["source_url"] and not public_link(record["source_url"]):reason.append("Source URL is invalid")
        if any(len(v)>MAX_CELL_LENGTH for v in record.values()):reason.append("Cell exceeds 30,000 characters")
        if len(record["subject"])>400 or "\n" in record["subject"] or "\r" in record["subject"]:reason.append("Invalid subject line")
        if any(len(v)>MAX_CELL_LENGTH for v in branch_data.values()):reason.append("Message branch exceeds 30,000 characters")
        # One Excel author may share a publishing website with another author;
        # name+country is an additional local file collision check.
        email_key="email:"+email if email else ""
        domain=(urlparse(website).hostname or "").lower().removeprefix("www.") if website else ""
        key=(email_key or "author:"+norm(person)+":"+norm(record["country"])) if service=="author" else ("site:"+domain if domain else email_key or "maps:"+place_id or "name:"+norm(person))
        if key in seen:reason.append("Duplicate in this file")
        if not reason:seen.add(key)
        # Uploaded old sending statuses are historical flags, not authorization
        # to mark new messages sent. Excluded statuses remain out of new drafts.
        status=norm(record["status"])
        suppressed=norm(record["do_not_contact"]) in TRUE_VALUES or any(status==v for v in EXCLUDED_STATUSES if v in ("opt out","unsubscribed","do not contact"))
        historical=status in EXCLUDED_STATUSES and not suppressed
        item={
            "row":number,"name":person,"email":email,"website":website,"maps_url":maps_url,
            "country":record["country"],"city":record["city"],
            "source_url":record["source_url"],"evidence_notes":record["evidence"],
            "subject":record["subject"],"first_message":msg,"first_message_english":record["english"],
            "branches":branch_data,"service_details":{k:v for k,v in record.items() if k in SERVICES[service] and k!="name" and v},
            "identity_key":key,"do_not_contact":suppressed,"historical_status":record["status"],
            "historical_only":historical,"errors":reason,
            "disposition":"invalid" if reason else "suppressed" if suppressed else "historical_skip" if historical else "ready_for_import",
        }
        (errors if reason else items).append(item)
    return {"sheet":name,"rows":items,"issues":errors,"columns":{k:str(header[i]) for k,i in columns.items()},"branch_columns":list(branches)}

def parse_service_upload(data,filename,service):
    if service not in SERVICES:raise ValueError("Unsupported Letter Desk service")
    if not isinstance(data,bytes) or len(data)>MAX_BYTES:raise ValueError("Workbook exceeds 64 MB.")
    lower=str(filename or "").lower()
    if lower.endswith(".csv"):
        try:
            reader=csv.reader(io.TextIOWrapper(io.BytesIO(data),encoding="utf-8-sig",newline=""))
            head=list(islice(reader,20))
            idx,cols,branches=_find_sheet(head,service)
            if idx is None:raise ValueError("No recognizable prospect headings in CSV")
            return _process_table("CSV",idx,head[idx],cols,branches,
                __import__('itertools').chain(head[idx+1:],reader),service)
        except (UnicodeDecodeError,csv.Error) as e:
            raise ValueError("Invalid CSV file") from e
    if not lower.endswith((".xlsx",".xlsm")):raise ValueError("Upload an .xlsx, .xlsm or .csv file.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos=z.infolist()
            if len(infos)>2000 or sum(x.file_size for x in infos)>MAX_UNCOMPRESSED:
                raise ValueError("Workbook exceeds safe expanded size")
            if any(x.flag_bits&1 for x in infos):raise ValueError("Password-protected workbook not supported")
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        try:
            for sheet in book.worksheets:
                if sheet.max_column and sheet.max_column>MAX_COLUMNS:continue
                first=list(islice(sheet.iter_rows(values_only=True),20))
                idx,cols,branches=_find_sheet(first,service)
                if idx is None:continue
                # If a spreadsheet contains >5000 rows, don't silently accept a partial import.
                if sheet.max_row and sheet.max_row>MAX_ROWS+idx+1:raise ValueError("Maximum 5,000 records per standard import; use Bulk upload for up to 50,000.")
                remaining=sheet.iter_rows(min_row=idx+2,values_only=True)
                return _process_table(sheet.title,idx,first[idx],cols,branches,remaining,service)
        finally:book.close()
    except ValueError:raise
    except Exception as e:raise ValueError("Invalid or unreadable Excel workbook") from e
    raise ValueError("No recognizable prospect headings in workbook")

def import_headers(service):
    if service not in SERVICES:raise ValueError("Unsupported Letter Desk service")
    base=["Name","Email","Website","Country","City","Source URL","Evidence Notes","Subject","First Message","First Message English","Do Not Contact"]
    fields=[aliases[0] for key,aliases in SERVICES[service].items() if key!="name"]
    return list(dict.fromkeys(base+fields+["Message 01","Message0101Yes","Message0101No","Message0101Question","Message0201Yes","Message0201No"]))


def iter_service_chunks(data,filename,service,chunk_size=500):
    """Yield (parsed partial table) without materializing all spreadsheet rows.

    Shared duplicate state and absolute row indices span all batches, even for
    XLSX. A trailing 50,001st row causes a hard failure before it is imported.
    """
    if service not in SERVICES:raise ValueError("Unsupported Letter Desk service")
    if not isinstance(data,bytes) or len(data)>MAX_BYTES:raise ValueError("Workbook exceeds 64 MB.")
    if type(chunk_size)!=int or not 1<=chunk_size<=1000:
        raise ValueError("Import chunk size must be 1–1,000 rows")
    filename=str(filename or "").lower()
    from itertools import chain
    def consume(name,idx,head,cols,branches,source):
        seen=set()
        count=0
        start=idx+2
        while True:
            section=list(islice(source,chunk_size))
            if not section:return
            # Blank spreadsheet formatting rows do not count as prospects.
            # Preserve absolute row positions even when there are blanks.
            valid_rows=sum(1 for row in section if any(clean_cell(v) for v in row))
            count+=valid_rows
            if count>MAX_BULK_ROWS:
                raise ValueError("Maximum 50,000 nonempty records per upload")
            part=_process_table(name,idx,head,cols,branches,section,service,
                                first_row=start,seen=seen)
            start+=len(section)
            if part["rows"] or part["issues"]:
                yield part
    if filename.endswith(".csv"):
        try:
            with io.BytesIO(data) as stream:
                textio=io.TextIOWrapper(stream,encoding="utf-8-sig",newline="")
                reader=csv.reader(textio)
                first=list(islice(reader,20))
                idx,cols,branches=_find_sheet(first,service)
                if idx is None:raise ValueError("No recognizable prospect headings in CSV")
                yield from consume("CSV",idx,first[idx],cols,branches,
                                   chain(first[idx+1:],reader))
        except (UnicodeDecodeError,csv.Error) as exc:
            raise ValueError("Invalid CSV file") from exc
        return
    if not filename.endswith((".xlsx",".xlsm")):
        raise ValueError("Upload an .xlsx, .xlsm or .csv file.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos=z.infolist()
            if (len(infos)>2000 or sum(x.file_size for x in infos)>MAX_UNCOMPRESSED
                or any(x.flag_bits&1 for x in infos)):
                raise ValueError("Workbook is too large or is password protected")
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        try:
            for sheet in book.worksheets:
                if sheet.max_column and sheet.max_column>MAX_COLUMNS:continue
                first=list(islice(sheet.iter_rows(values_only=True),20))
                idx,cols,branches=_find_sheet(first,service)
                if idx is None:continue
                yield from consume(sheet.title,idx,first[idx],cols,branches,
                    sheet.iter_rows(min_row=idx+2,values_only=True))
                return
        finally:
            book.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Invalid or unreadable Excel workbook") from exc
    raise ValueError("No recognizable prospect headings in workbook")
